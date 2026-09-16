import csv
import json
import os

import fitz  # PyMuPDF
import pymupdf4llm
from openpyxl import load_workbook

from RagCore.Ingestion.exceptions import EmptyExtractionError, IngestionError
from RagCore.Ingestion.document import (
    DOCUMENT_LAYOUT_FORMATS, 
    DOCUMENT_TEXT_FORMATS,
    ALL_ALLOWED_EXTENSIONS,
    DocumentPage,
)

class DocumentExtractionEngine:

    @classmethod
    def extract_to_markdown(cls, file_path: str) -> list[DocumentPage]:
        if not os.path.exists(file_path):
            raise IngestionError(f"Target document file not found at: {file_path}")

        _, ext = os.path.splitext(file_path.lower())

        # 1. Global Guardrail
        if ext not in ALL_ALLOWED_EXTENSIONS:
            raise IngestionError(f"File extension '{ext}' is not supported by this system.")

        try:
            # 2. Handle Complex Layout Formats (.pdf, .epub, .html, etc.) via PyMuPDF4LLM
            if ext in DOCUMENT_LAYOUT_FORMATS:
                page_dicts = pymupdf4llm.to_markdown(file_path, page_chunks=True)
                pages = [
                    DocumentPage(
                        page_number=page_dict["metadata"]["page_number"],
                        content=str(page_dict["text"]),
                    )
                    for page_dict in page_dicts
                ]

                if not any(page.content.strip() for page in pages):
                    raise EmptyExtractionError(
                        "Document text space came up empty or is corrupted."
                    )

                return pages

            # 3. Handle standard text/data formats.
            elif ext in DOCUMENT_TEXT_FORMATS:
                if ext in {".txt", ".md"}:
                    with open(file_path, encoding="utf-8-sig") as source:
                        content = source.read()
                elif ext == ".csv":
                    with open(file_path, newline="", encoding="utf-8-sig") as source:
                        content = "\n".join(
                            " | ".join(row) for row in csv.reader(source)
                        )
                elif ext == ".json":
                    with open(file_path, encoding="utf-8-sig") as source:
                        content = json.dumps(json.load(source), indent=2, ensure_ascii=False)
                elif ext == ".xlsx":
                    workbook = load_workbook(file_path, read_only=True, data_only=True)
                    try:
                        rows = []
                        for worksheet in workbook.worksheets:
                            rows.append(f"# {worksheet.title}")
                            rows.extend(
                                " | ".join("" if value is None else str(value) for value in row)
                                for row in worksheet.iter_rows(values_only=True)
                            )
                        content = "\n".join(rows)
                    finally:
                        workbook.close()
                else:
                    with fitz.open(file_path) as doc:
                        content = "\n".join(page.get_text("text") for page in doc)

                pages = [DocumentPage(page_number=1, content=content)]

                if not content.strip():
                    raise EmptyExtractionError(
                        "Document text space came up empty or is corrupted."
                    )

                return pages

            else:
                raise IngestionError(f"Extension '{ext}' must be routed to the Vision Engine.")

        # FIX: Complete block adding except clause to resolve try statement error
        except Exception as exc:
            if isinstance(exc, IngestionError):
                raise exc
            raise IngestionError(
                f"Failed parsing file via Core Ingest Engine: {str(exc)}"
            ) from exc